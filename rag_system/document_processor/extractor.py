# document_processor/extractor.py
# ----------------------------------------------------------
# Extracts raw text from PDF, DOCX, and PPTX files.
# Each extracted piece carries its page number and source
# file name so we can trace every chunk back to its origin.
# ----------------------------------------------------------

import os
from dataclasses import dataclass, field
from typing import List

from loguru import logger


# ── Data structure returned by every extractor ───────────
@dataclass
class PageContent:
    """
    Holds text from a single page (or slide) of a document.

    Fields:
        source_file  : original filename (e.g. 'lecture1.pdf')
        page_number  : 1-based page / slide number
        text         : raw text from that page
    """
    source_file: str
    page_number: int
    text: str


# ── PDF ───────────────────────────────────────────────────
def extract_pdf(file_path: str) -> List[PageContent]:
    """
    Extract text from every page of a PDF.
    Uses PyMuPDF (fitz) which handles complex layouts well.
    """
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ImportError("Install PyMuPDF:  pip install PyMuPDF")

    pages = []
    file_name = os.path.basename(file_path)

    with fitz.open(file_path) as doc:
        total_pages = len(doc)
        logger.info(f"PDF '{file_name}' has {total_pages} pages.")

        for page_index in range(total_pages):
            page = doc[page_index]
            text = page.get_text("text")   # plain text extraction

            if text.strip():               # skip completely blank pages
                pages.append(PageContent(
                    source_file=file_name,
                    page_number=page_index + 1,   # convert 0-based → 1-based
                    text=text.strip()
                ))
            else:
                logger.debug(f"  Skipping blank page {page_index + 1}")

    logger.success(f"Extracted {len(pages)} non-blank pages from '{file_name}'.")
    return pages


# ── DOCX ──────────────────────────────────────────────────
def extract_docx(file_path: str) -> List[PageContent]:
    """
    Extract text from a Word document.
    Word documents don't have hard page boundaries, so we treat
    every 50 non-empty paragraphs as one logical 'page'.
    This keeps chunk sizes manageable.
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError("Install python-docx:  pip install python-docx")

    file_name = os.path.basename(file_path)
    doc = Document(file_path)

    # Collect all non-empty paragraph texts
    all_paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    logger.info(f"DOCX '{file_name}' has {len(all_paragraphs)} non-empty paragraphs.")

    pages = []
    paragraphs_per_page = 50   # treat every 50 paragraphs as one logical page
    page_number = 1

    for i in range(0, len(all_paragraphs), paragraphs_per_page):
        chunk_paragraphs = all_paragraphs[i : i + paragraphs_per_page]
        page_text = "\n".join(chunk_paragraphs)

        pages.append(PageContent(
            source_file=file_name,
            page_number=page_number,
            text=page_text
        ))
        page_number += 1

    logger.success(f"Extracted {len(pages)} logical pages from '{file_name}'.")
    return pages


# ── PPTX ──────────────────────────────────────────────────
def extract_pptx(file_path: str) -> List[PageContent]:
    """
    Extract text from a PowerPoint presentation.
    Each slide becomes one 'page'.
    """
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("Install python-pptx:  pip install python-pptx")

    file_name = os.path.basename(file_path)
    prs = Presentation(file_path)
    pages = []

    logger.info(f"PPTX '{file_name}' has {len(prs.slides)} slides.")

    for slide_index, slide in enumerate(prs.slides):
        slide_texts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())

        full_text = "\n".join(slide_texts)

        if full_text.strip():
            pages.append(PageContent(
                source_file=file_name,
                page_number=slide_index + 1,   # 1-based slide number
                text=full_text
            ))
        else:
            logger.debug(f"  Skipping blank slide {slide_index + 1}")

    logger.success(f"Extracted {len(pages)} non-blank slides from '{file_name}'.")
    return pages


# ── Main dispatcher ───────────────────────────────────────
def extract_document(file_path: str) -> List[PageContent]:
    """
    Auto-detects file type and calls the right extractor.
    Supports: .pdf, .docx, .pptx

    Args:
        file_path : full path to the uploaded document

    Returns:
        List of PageContent objects, one per page/slide.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return extract_pdf(file_path)
    elif ext == ".docx":
        return extract_docx(file_path)
    elif ext == ".pptx":
        return extract_pptx(file_path)
    else:
        raise ValueError(
            f"Unsupported file type: '{ext}'. "
            f"Supported types: .pdf, .docx, .pptx"
        )